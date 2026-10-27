from __future__ import annotations
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from .common import ROOT, OUTPUT_DIR

NAVY=RGBColor(21,34,56); TEAL=RGBColor(23,107,135); PALE=RGBColor(241,246,250); INK=RGBColor(38,52,69); MUTED=RGBColor(83,98,118); RED=RGBColor(150,47,47); WHITE=RGBColor(255,255,255)

def add_text(slide, text, x,y,w,h, size=18, color=INK, bold=False, font='Aptos', align=None, margin=0.05):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin)
    p=tf.paragraphs[0]; p.text=str(text); p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
    if align is not None: p.alignment=align
    return box

def add_bg(slide):
    fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=WHITE
    strip=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0, Inches(0.18), Inches(7.5))
    strip.fill.solid(); strip.fill.fore_color.rgb=TEAL; strip.line.fill.background()
    foot=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.18), Inches(7.33), Inches(12.32), Inches(0.17))
    foot.fill.solid(); foot.fill.fore_color.rgb=NAVY; foot.line.fill.background()

def add_title(slide,title,number):
    add_text(slide,title,0.55,0.25,12.1,0.75,size=24,color=NAVY,bold=True)
    add_text(slide,f'SOSYAL AĞ ANALİZİ  •  2025 VERİSİ  •  {number}/9',0.57,7.02,11.8,0.18,size=8,color=MUTED)

def add_bullets(slide, bullets, x=0.8,y=1.4,w=11.4,h=4.9, size=18, color=INK):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(0.05); tf.margin_right=Inches(0.05); tf.margin_top=Inches(0.05)
    for i,item in enumerate(bullets):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text='• '+str(item); p.font.name='Aptos'; p.font.size=Pt(size); p.font.color.rgb=color; p.space_after=Pt(13)
    return box

def add_card(slide,x,y,w,h,big,label,color=TEAL):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=PALE; sh.line.color.rgb=RGBColor(218,226,235)
    add_text(slide,big,x+0.12,y+0.12,w-0.24,0.55,size=26,color=color,bold=True,align=PP_ALIGN.CENTER)
    add_text(slide,label,x+0.08,y+0.75,w-0.16,h-0.8,size=11,color=MUTED,align=PP_ALIGN.CENTER)

def main():
    out=OUTPUT_DIR/'sunum_TASLAK.pptx'
    summary=json.loads((OUTPUT_DIR/'analysis_summary.json').read_text(encoding='utf-8'))
    audit=json.loads((OUTPUT_DIR/'data_audit.json').read_text(encoding='utf-8'))
    g=summary['graph']; comm=summary['community_detection']; null=summary['null_model']; closure=summary['closure_scenario']
    measures=__import__('pandas').read_csv(OUTPUT_DIR/'centrality_and_communities.csv')
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    blank=prs.slide_layouts[6]
    # Slide 1
    s=prs.slides.add_slide(blank); add_bg(s)
    add_text(s,'İSTANBUL RAYLI SİSTEM YOLCU AĞI',0.7,1.0,11.7,0.8,size=30,color=NAVY,bold=True)
    add_text(s,'İstasyon merkeziyeti, yolcu aktivitesi ve istasyon kapanma senaryosu',0.75,2.0,11.1,0.6,size=20,color=TEAL)
    add_text(s,'Sosyal Ağ Analizi • Oğuzhan Sapmaz • 2516016021',0.77,3.0,10.5,0.35,size=13,color=MUTED)
    warn=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(4.0), Inches(11.5), Inches(0.9)); warn.fill.solid(); warn.fill.fore_color.rgb=RGBColor(255,245,245); warn.line.color.rgb=RGBColor(230,184,184)
    add_text(s,'Önemli sınırlar: OD akışı yok; bazı isim/aktarma eşleştirmeleri ve 500 düğüm şartı ayrıca ele alınmalı.',0.95,4.2,11.05,0.55,size=14,color=RED,bold=True)
    # Slide 2
    s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'2025 CSV 135.783 kayıt ve 23 hat içeriyor.',2)
    add_card(s,0.7,1.4,2.6,1.5,f"{audit['rows']:,}",'Kaynak kayıt',TEAL); add_card(s,3.55,1.4,2.6,1.5,audit['distinct_lines'],'Hat',TEAL); add_card(s,6.4,1.4,2.6,1.5,audit['distinct_station_numbers'],'Farklı istasyon kodu',TEAL); add_card(s,9.25,1.4,2.6,1.5,audit['columns'],'Kaynak sütun',TEAL)
    add_bullets(s,[f"Yıl: {', '.join(audit['years'])}; 12 ayın tamamında kayıt bulunuyor.",f"{audit['missing_station_name_records']:,} kayıtta istasyon adı boş.",f"{audit['duplicate_full_source_rows']} tam yinelenen kaynak satırı tespit edildi."],y=3.5,size=16)
    # Slide 3
    s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'Koordinatlar analizden önce biçim ve aralık kontrolünden geçirildi.',3)
    add_bullets(s,[f"Kaynakta longitude/latitude alanları gruplu noktalar nedeniyle standart ondalık biçimde değil.",f"Dönüşüm sonrası İstanbul aralığında {audit['normalized_valid_coordinate_records']:,} kayıtta koordinat elde edildi.",f"Her koordinat ekseninde {audit['missing_longitude_source_records']:,} kaynak kaydı boş.","Eksik koordinatlar tahmin edilmedi; düğüm konumları şimdilik ağ yerleşimiyle görselleştiriliyor."],y=1.45,size=17)
    # Slide 4
    s=prs.slides.add_slide(blank); add_bg(s); add_title(s,f"Düzeltilmiş topoloji {g['nodes']} düğüm ve {g['edges']} kenar içeriyor.",4)
    img=OUTPUT_DIR/'plots'/'network_overview.png'
    if img.exists(): s.shapes.add_picture(str(img), Inches(0.55), Inches(1.25), width=Inches(7.2))
    add_bullets(s,[f"Bağlı bileşen sayısı: {g['connected_components']}.",f"En büyük bileşen: {g['largest_component_nodes']} düğüm / {g['largest_component_edges']} kenar.","Kenarlar, aynı hattaki ardışık istasyonları temsil ediyor.","Sıra sayısal olarak işlenir; docs/topology_validation.csv hat kaynaklarını ve kalan kontrolleri listeler."],x=7.95,y=1.55,w=4.6,h=4.5,size=15)
    # Slide 5
    s=prs.slides.add_slide(blank); add_bg(s); add_title(s,f"En büyük bağlı bileşende ortalama yol uzunluğu {fmt(g['average_shortest_path_largest_component'],2)} ve çap {g['diameter_largest_component']}.",5)
    add_card(s,0.8,1.45,2.65,1.55,g['largest_component_nodes'],'En büyük bileşen düğümü'); add_card(s,3.75,1.45,2.65,1.55,fmt(g['density']),'Yoğunluk'); add_card(s,6.7,1.45,2.65,1.55,fmt(g['average_clustering_all_graph'],4),'Ortalama kümelenme'); add_card(s,9.65,1.45,2.65,1.55,fmt(g['degree_assortativity'],4),'Derece assortativity')
    add_bullets(s,["Ağ büyük ölçüde hat zincirleri ve aktarma noktalarından oluşuyor.","Üçgen olmadığı için ortalama kümelenme sıfır çıktı; bu, bu ön topoloji modelinin bir özelliği ve doğrulanmalı.","Ortalama yol ve çap yalnızca en büyük bağlı bileşen için hesaplandı."],y=3.55,size=16)
    # Slide 6
    s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'Merkeziyet sıralaması topolojik aracı istasyonları öne çıkarıyor.',6)
    top=summary['top_centrality_tables'].get('betweenness_unweighted',[])[:8]
    rows=[['İstasyon düğümü','Betweenness']]+[[r['label'],f"{r['betweenness_unweighted']:.5f}"] for r in top]
    table=s.shapes.add_table(len(rows),2, Inches(0.75), Inches(1.35), Inches(6.3), Inches(4.75)).table
    table.columns[0].width=Inches(4.5); table.columns[1].width=Inches(1.8)
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            cell=table.cell(i,j); cell.text=str(val); cell.margin_left=Inches(0.08); cell.margin_right=Inches(0.05)
            cell.fill.solid(); cell.fill.fore_color.rgb= NAVY if i==0 else (PALE if i%2 else WHITE)
            for p in cell.text_frame.paragraphs:
                p.font.name='Aptos'; p.font.size=Pt(11); p.font.bold=(i==0); p.font.color.rgb=WHITE if i==0 else INK
    add_bullets(s,["Ağırlıksız betweenness, hatlar arası kısa yolların sık geçtiği düğümleri ölçer.","Ağırlıklı ölçüde kenar mesafesi, uç istasyonların ortalama günlük yolcu aktivitesinin tersi olarak tanımlandı.","İki sıralama arasındaki fark araştırma sorusuna cevap verir; sonuçlar topoloji doğrulamasına duyarlı."],x=7.4,y=1.6,w=5.1,h=4.5,size=14)
    # Slide 7
    s=prs.slides.add_slide(blank); add_bg(s); add_title(s,f"Louvain ön çalışması {comm['louvain']['community_count']} topluluk buldu.",7)
    add_card(s,0.9,1.45,3.0,1.65,comm['louvain']['community_count'],'Louvain topluluk sayısı'); add_card(s,4.2,1.45,3.0,1.65,fmt(comm['louvain']['modularity'],4),'Modülerlik Q'); add_card(s,7.5,1.45,3.0,1.65,fmt(summary['community_detection'].get('louvain_stability',{}).get('mean_ARI_vs_seed42'),3),'Ortalama ARI (10 tohum)')
    topcom=summary.get('louvain_community_details',[])[:4]
    add_bullets(s,[f"{c.get('candidate_name','Topluluk')}: {c['size']} düğüm; baskın hatlar {', '.join(c.get('dominant_lines',[]))}; yolcu aktivitesi yüksek istasyonlar: {', '.join(c['top_passenger_stations'][:2]) or 'eşleşmiş veri yok'}." for c in topcom],y=3.45,size=13)
    # Slide 8
    s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'İstasyon kapanma testi gerçek yolcu kayması değil, topolojik yeniden yönelmedir.',8)
    add_text(s,f"Kapanma adayı: {closure.get('station','hesaplanmadı')}",0.85,1.35,11.3,0.45,size=20,color=TEAL,bold=True)
    add_bullets(s,[f"Bileşen sayısı: {closure.get('components_before','?')} → {closure.get('components_after','?')}.",f"En büyük bileşen düğümü: {closure.get('largest_component_nodes_before','?')} → {closure.get('largest_component_nodes_after','?')}.",f"Bağlı düğüm çifti oranı: {fmt(closure.get('connected_pair_fraction_before'),3)} → {fmt(closure.get('connected_pair_fraction_after'),3)}; global efficiency: {fmt(closure.get('global_efficiency_before'),4)} → {fmt(closure.get('global_efficiency_after'),4)}.","Bu, düğüm ve kenarlarının silindiği idealize topolojik senaryodur; gerçek yolcu kayması değildir.","OD matrisi olmadığı için yolcu hacminin hangi hatta ne kadar kayacağı hesaplanamaz."],y=2.1,size=15)
    # Slide 9
    s=prs.slides.add_slide(blank); add_bg(s); add_title(s,'Sonuçlar ve teslim öncesi açık kontroller.',9)
    nulltxt=(f"Null model: {null.get('replicates_completed',0)}/{null.get('replicates_requested',0)} tekrar" if null.get('status')=='calculated' else 'Null model durum: '+str(null.get('status')))
    leiden_status=str(comm.get('leiden',{}).get('status','bilinmiyor'))
    powerlaw_status=str(summary.get('powerlaw',{}).get('status','bilinmiyor'))
    add_bullets(s,[nulltxt,'Leiden: '+leiden_status,'Power-law / exponential / lognormal karşılaştırması: '+powerlaw_status,'Genel yönergede en az 500 düğüm şartı var; mevcut gerçek istasyon grafiği daha küçük. Düğüm uydurulmadı; kapsam öğretim elemanıyla netleştirilmeli.','İstasyonlar arası OD matrisi bulunmadığından kenar ağırlığı yolcu aktivitesi vekilidir; kapanma analizi gerçek yolcu yük kayması değildir.','Etkileşimli HTML: web/interactive_network.html; A3 poster: outputs/ag_posteri_A3_TASLAK.pdf.'],y=1.4,size=14,color=INK)
    warn=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(6.05), Inches(11.4), Inches(0.65)); warn.fill.solid(); warn.fill.fore_color.rgb=RGBColor(255,245,245); warn.line.color.rgb=RGBColor(230,184,184)
    add_text(s,'Sunumda her yöntemi kendi cümlenle açıkla; hesaplanmayan veya doğrulanmayan sonucu kesin bulgu olarak sunma.',1.05,6.19,11.0,0.35,size=12,color=RED,bold=True)
    prs.save(out); print('Presentation created:',out)

def fmt(x,d=3):
    try: return f'{float(x):.{d}f}'
    except Exception: return str(x)

if __name__=='__main__': main()
