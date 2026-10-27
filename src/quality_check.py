from __future__ import annotations
import json
import re
from pathlib import Path
import pandas as pd
import networkx as nx
from pptx import Presentation

from .common import ROOT, OUTPUT_DIR, ensure_dirs, write_json


def main() -> int:
    ensure_dirs()
    checks = []
    warnings = []
    def check(name, ok, detail=''):
        checks.append({'check': name, 'status': 'PASS' if ok else 'FAIL', 'detail': str(detail)})
    def warn(name, detail):
        warnings.append({'item': name, 'detail': str(detail)})

    raw_files = list((ROOT/'data'/'raw').glob('*.csv'))
    check('Raw CSV exists', bool(raw_files), ', '.join(p.name for p in raw_files))
    required = [
        OUTPUT_DIR/'istanbul_rail_2025.graphml', OUTPUT_DIR/'nodes.csv', OUTPUT_DIR/'edges.csv',
        OUTPUT_DIR/'centrality_and_communities.csv', OUTPUT_DIR/'null_model_results.csv',
        OUTPUT_DIR/'analysis_summary.json', OUTPUT_DIR/'rapor_TASLAK.pdf', OUTPUT_DIR/'ag_posteri_A3_TASLAK.pdf',
        OUTPUT_DIR/'sunum_TASLAK.pptx', OUTPUT_DIR/'sunum_TASLAK.pdf',
        ROOT/'web'/'interactive_network.html', ROOT/'notebooks'/'01_network_analysis.ipynb',
        ROOT/'topology'/'line_stations.csv', ROOT/'topology'/'transfer_edges.csv',
        ROOT/'docs'/'data_dictionary.csv', ROOT/'docs'/'topology_validation.csv',
    ]
    check('Required deliverable files exist', all(p.exists() and p.stat().st_size > 0 for p in required),
          f'{sum(p.exists() and p.stat().st_size > 0 for p in required)}/{len(required)} files present')

    summary_path = OUTPUT_DIR/'analysis_summary.json'
    summary = json.loads(summary_path.read_text(encoding='utf-8')) if summary_path.exists() else {}
    graph_path = OUTPUT_DIR/'istanbul_rail_2025.graphml'
    G = nx.read_graphml(graph_path) if graph_path.exists() else nx.Graph()
    metadata = json.loads((OUTPUT_DIR/'graph_metadata.json').read_text(encoding='utf-8')) if (OUTPUT_DIR/'graph_metadata.json').exists() else {}
    nodes = pd.read_csv(OUTPUT_DIR/'nodes.csv', encoding='utf-8-sig') if (OUTPUT_DIR/'nodes.csv').exists() else pd.DataFrame()
    edges = pd.read_csv(OUTPUT_DIR/'edges.csv', encoding='utf-8-sig') if (OUTPUT_DIR/'edges.csv').exists() else pd.DataFrame()
    check('GraphML node count matches summary', G.number_of_nodes() == summary.get('graph', {}).get('nodes'), G.number_of_nodes())
    check('GraphML edge count matches summary', G.number_of_edges() == summary.get('graph', {}).get('edges'), G.number_of_edges())
    check('GraphML matches node/edge CSVs', len(nodes)==G.number_of_nodes() and len(edges)==G.number_of_edges(), f'nodes={len(nodes)}, edges={len(edges)}')
    check('Graph has no self-loops', nx.number_of_selfloops(G)==0, nx.number_of_selfloops(G))
    check('Every edge endpoint exists', all(u in G and v in G for u,v in G.edges()), 'GraphML consistency')

    topo = pd.read_csv(ROOT/'topology'/'line_stations.csv', dtype={'station_order':str}, encoding='utf-8-sig')
    topo['station_order_num'] = pd.to_numeric(topo['station_order'], errors='coerce')
    monotonic = True
    duplicate_orders = []
    for route, group in topo.groupby('route_id'):
        vals = group['station_order_num'].tolist()
        if vals != sorted(vals) or len(vals) != len(set(vals)) or any(pd.isna(vals)):
            monotonic = False
            duplicate_orders.append(route)
    check('Topology order is numeric, sorted and unique per route', monotonic, 'problem routes: '+', '.join(duplicate_orders))
    transfer = pd.read_csv(ROOT/'topology'/'transfer_edges.csv', encoding='utf-8-sig')
    check('Transfer edge endpoints are in graph', all(str(r.node_u) in G and str(r.node_v) in G for r in transfer.itertuples()), f'{len(transfer)} explicit transfer edges')
    check('Transfer edges exist in GraphML', all(G.has_edge(str(r.node_u), str(r.node_v)) for r in transfer.itertuples()), f'{len(transfer)} transfer links')

    centrality = pd.read_csv(OUTPUT_DIR/'centrality_and_communities.csv', encoding='utf-8-sig') if (OUTPUT_DIR/'centrality_and_communities.csv').exists() else pd.DataFrame()
    check('Centrality rows match nodes', len(centrality)==G.number_of_nodes(), len(centrality))
    core_cols = [c for c in ['degree','betweenness_unweighted','betweenness_weighted','current_flow_betweenness_lcc'] if c in centrality.columns]
    core_ok = bool(core_cols) and not centrality[core_cols].isna().any().any()
    check('Core centrality values contain no missing cells', core_ok, ', '.join(core_cols))

    null_path = OUTPUT_DIR/'null_model_results.csv'
    null_df = pd.read_csv(null_path, encoding='utf-8-sig') if null_path.exists() else pd.DataFrame()
    null = summary.get('null_model', {})
    check('100 null model samples completed', len(null_df)==100 and null.get('replicates_completed')==100, f'rows={len(null_df)}, status={null.get("status")}')
    if len(null_df):
        if 'is_connected' in null_df.columns:
            check('All null graphs connected', bool(null_df['is_connected'].all()), int(null_df['is_connected'].sum()))
        else: check('All null graphs connected', False, 'is_connected column missing')
        if 'degree_sequence_preserved' in null_df.columns:
            check('Degree sequence preserved in every null graph', bool(null_df['degree_sequence_preserved'].all()), int(null_df['degree_sequence_preserved'].sum()))
        else: check('Degree sequence preserved in every null graph', False, 'degree_sequence_preserved column missing')
        if 'edge_set_hash' in null_df.columns:
            check('Null graphs are unique', null_df['edge_set_hash'].nunique()==len(null_df), f'unique={null_df["edge_set_hash"].nunique()}')
        else: check('Null graphs are unique', False, 'edge_set_hash column missing')

    notebook_path = ROOT/'notebooks'/'01_network_analysis.ipynb'
    try:
        # Use the standard library to validate the notebook so QA does not
        # require the optional nbformat package to be installed.
        nb_data = json.loads(notebook_path.read_text(encoding='utf-8'))
        cells = nb_data.get('cells', [])
        if not isinstance(cells, list):
            raise ValueError("Notebook 'cells' alanı bir liste değil.")
        if not str(nb_data.get('nbformat', '')).isdigit() or int(nb_data.get('nbformat', 0)) < 4:
            raise ValueError("Notebook nbformat sürümü 4 veya üzeri değil.")
        code_count = 0
        for i, cell in enumerate(cells):
            if not isinstance(cell, dict):
                raise ValueError(f'{i}. hücre nesne değil.')
            if cell.get('cell_type') == 'code':
                source = cell.get('source', '')
                if isinstance(source, list):
                    source = ''.join(source)
                if not isinstance(source, str):
                    raise ValueError(f'{i}. hücrenin source alanı metin değil.')
                compile(source, f'notebook_cell_{i}', 'exec')
                code_count += 1
        check('Notebook parses and code cells compile', True, f'{len(cells)} cells; {code_count} code cells; stdlib JSON validation')
    except Exception as exc:
        check('Notebook parses and code cells compile', False, exc)

    prs = Presentation(OUTPUT_DIR/'sunum_TASLAK.pptx') if (OUTPUT_DIR/'sunum_TASLAK.pptx').exists() else None
    check('PowerPoint has 8–10 slides', prs is not None and 8 <= len(prs.slides) <= 10, len(prs.slides) if prs else 'missing')
    try:
        try:
            # Use pypdf when available.
            from pypdf import PdfReader
            report_pages = len(PdfReader(str(OUTPUT_DIR/'rapor_TASLAK.pdf')).pages)
            slide_pdf_pages = len(PdfReader(str(OUTPUT_DIR/'sunum_TASLAK.pdf')).pages)
            poster_reader = PdfReader(str(OUTPUT_DIR/'ag_posteri_A3_TASLAK.pdf'))
            poster_page = poster_reader.pages[0]
            width_pt = float(poster_page.mediabox.width)
            height_pt = float(poster_page.mediabox.height)
            pdf_reader_note = 'pypdf'
        except ImportError:
            # Fallback for environments that do not have the optional pypdf package.
            # The project's generated PDFs keep page objects and MediaBox values readable.
            def pdf_page_details(path):
                raw = path.read_bytes()
                page_count = len(re.findall(rb'/Type\s*/Page\b', raw))
                boxes = re.findall(
                    rb'/MediaBox\s*\[\s*([-+\d.]+)\s+([-+\d.]+)\s+([-+\d.]+)\s+([-+\d.]+)\s*\]',
                    raw,
                )
                if page_count < 1 or not boxes:
                    raise ValueError(f'PDF sayfa bilgisi okunamadı: {path.name}; pypdf kurulu değil.')
                parsed = tuple(float(x) for x in boxes[0])
                return page_count, abs(parsed[2] - parsed[0]), abs(parsed[3] - parsed[1])
            report_pages, _, _ = pdf_page_details(OUTPUT_DIR/'rapor_TASLAK.pdf')
            slide_pdf_pages, _, _ = pdf_page_details(OUTPUT_DIR/'sunum_TASLAK.pdf')
            _, width_pt, height_pt = pdf_page_details(OUTPUT_DIR/'ag_posteri_A3_TASLAK.pdf')
            pdf_reader_note = 'standart kütüphane PDF işaretleri (pypdf yok)'
        is_a3 = (abs(width_pt-1190.55)<3 and abs(height_pt-841.89)<3) or (abs(width_pt-841.89)<3 and abs(height_pt-1190.55)<3)
        check('Report PDF has 6–10 pages', 6 <= report_pages <= 10, f'{report_pages} pages; {pdf_reader_note}')
        check('Presentation PDF has 8–10 pages', 8 <= slide_pdf_pages <= 10, f'{slide_pdf_pages} pages; {pdf_reader_note}')
        check('Poster PDF is A3', is_a3, f'{width_pt:.1f} x {height_pt:.1f} pt; {pdf_reader_note}')
    except Exception as exc:
        check('PDF dimensions/page counts', False, exc)

    leiden = summary.get('community_detection', {}).get('leiden', {})
    powerlaw = summary.get('powerlaw', {})
    if leiden.get('status') != 'calculated': warn('Leiden', leiden.get('status', 'not calculated'))
    if powerlaw.get('status') != 'calculated': warn('Power-law comparisons', powerlaw.get('reason', powerlaw.get('status','not calculated')))
    if G.number_of_nodes() < 500: warn('500-node rubric', f'Current real station graph has {G.number_of_nodes()} nodes; assignment card does not state an exception. Do not fabricate nodes; clarify scope with instructor.')
    warn('OD matrix limitation', 'Source contains station-level counts but no origin-destination trips. Edge strength is a passenger-activity proxy, not observed inter-station flow.')
    warn('Topology review', 'Official route pages were spot-checked and explicit transfer links added, but the full line/alias/transfer list still needs human review.')
    warn('License review', 'The dataset licence/attribution URL is recorded in the fetch script, but the exact terms were not independently confirmed in this run.')
    warn('Portfolio deployment', 'Interactive HTML must still be tested in a browser and embedded/published in the student portfolio.')
    pptx_path = OUTPUT_DIR/'sunum_TASLAK.pptx'
    slide_pdf_path = OUTPUT_DIR/'sunum_TASLAK.pdf'
    if pptx_path.exists() and slide_pdf_path.exists() and slide_pdf_path.stat().st_mtime < pptx_path.stat().st_mtime - 2:
        warn('Presentation PDF freshness', 'Sunum PDF dosyası PPTX dosyasından eski. PowerShell dışa aktarımı bu bilgisayarda engellendiyse PowerPoint içinden File > Export / Save As > PDF ile yeniden dışa aktar.')

    result = {
        'audit_timestamp': '2026-10-27',
        'status': 'PASS_WITH_WARNINGS' if all(x['status']=='PASS' for x in checks) else 'FAIL',
        'checks': checks,
        'warnings': warnings,
        'graph': {'nodes': G.number_of_nodes(), 'edges': G.number_of_edges(), 'components': sorted((len(c) for c in nx.connected_components(G)), reverse=True)},
        'analysis_status': {'leiden': leiden.get('status','not calculated'), 'powerlaw': powerlaw.get('status','not calculated'), 'null_model': null.get('status','not calculated')},
    }
    write_json(OUTPUT_DIR/'quality_check.json', result)
    lines=['SCIENTIFIC / FILE QUALITY CHECK', f"Overall: {result['status']}", '']
    lines += [f"[{x['status']}] {x['check']}: {x['detail']}" for x in checks]
    lines += ['', 'WARNINGS / NOT FULLY AUTOMATABLE']
    lines += [f"[!] {x['item']}: {x['detail']}" for x in warnings]
    (OUTPUT_DIR/'quality_check.txt').write_text('\n'.join(lines)+'\n', encoding='utf-8')
    print('\n'.join(lines))
    return 0 if result['status'] != 'FAIL' else 1

if __name__ == '__main__':
    raise SystemExit(main())
